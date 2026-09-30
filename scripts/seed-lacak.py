"""Seeder: sheet LACAK (LACAK KONSUMEN.xlsx) -> Supabase tabel lacak_konsumen (+ salinan data/lacak-konsumen.json)

Ambil seksi DATA PEMESAN + PRABOOKING, TEKNIK, LEGAL, KEUANGAN.
Likuiditas belum diambil.

    python3 scripts/seed-lacak.py [path.xlsx]

Butuh SUPABASE_URL + SUPABASE_SECRET_KEY (atau SUPABASE_SERVICE_ROLE_KEY) di env atau .env.local.
Excel = sumber kebenaran: baris yang hilang dari Excel ikut dihapus dari tabel.
"""
# ponytail: python + openpyxl (sudah ada di mesin) supaya tidak menambah dependency xlsx ke app
import json
import os
import sys
import urllib.parse
import urllib.request
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "LACAK KONSUMEN.xlsx"
OUT = ROOT / "data" / "lacak-konsumen.json"
FIRST_ROW = 14  # baris 11-13 = header bertingkat
# ponytail: Excel tidak punya kolom NIK; NIK dummy per blok untuk uji lookup, hapus saat NIK asli tersedia
NIK_DUMMY = {"B-3": "3329120806900001"}

PEMESANAN = {
    "no": "C", "id_unit": "D", "status_unit": "E", "nama": "F", "jenis_unit": "G", "blok": "H",
    "tipe": "I", "luas": "J", "harga": "K", "tgl_order": "L", "sales": "M", "sales_lengkap": "N",
    "divisi": "O", "sumber": "P", "referal": "Q", "metode_pembelian": "R", "bank": "S",
    "status_sop_sales": "V", "tgl_status_sop_sales": "W",
    "slik": "X", "spu": "Y", "spjb": "Z", "profiling": "AA", "go_to_likuid": "AB",
}
TAHAP = ["pondasi", "plat_2_lantai", "atap", "fasad", "hitaman", "bast", "rumah_selesai"]
# unit indent vs ready stock punya blok kolom sendiri, strukturnya sama
TEKNIK_INDENT = dict(spmk="BK", lantai="BL", progress="BM", **dict(zip(TAHAP, ["BN", "BO", "BP", "BQ", "BR", "BS", "BT"])))
TEKNIK_READY = dict(spmk="BU", lantai="BV", progress="BW", **dict(zip(TAHAP, ["BX", "BY", "BZ", "CA", "CB", "CC", "CD"])))
LEGAL = {
    "pra_akad": "CF", "go_to_legal": "CG", "akad": "CH", "shm": "CI",
    "notaris_pemberkasan": "CJ", "notaris_titip_ttd": "CK", "ajb_pph": "CL", "ajb_bphtb": "CM",
    "bpn_bn": "CN", "bpn_ph": "CO",
}
KEUANGAN = {
    "status_unit": "CR", "tum": "CS", "booking": "CT", "um": "CU", "ct": "CV",
    "piutang_um": "CW", "piutang_ct": "CX", "pencairan_kpr": "CY", "akad_kpr": "CZ",
    "dt1_pondasi": "DA", "dt2_atap": "DB", "dt3_bast": "DC", "dt4_lainnya": "DD",
    "penjualan": "DE", "dana_masuk": "DF", "status_penjualan": "DG", "tgl_penjualan": "DH", "pph": "DI",
}


def clean(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        v = v.strip()
        return None if v in ("", "-", "#N/A", "end") or v.startswith("Ref_") else v
    return v


def pick(row, cols):
    return {k: clean(row[column_index_from_string(c) - 1]) for k, c in cols.items()}


def teknik(row):
    t = pick(row, TEKNIK_INDENT)
    jenis = "indent"
    if t["spmk"] is None and t["progress"] is None:
        ready = pick(row, TEKNIK_READY)
        if ready["spmk"] or ready["progress"] is not None:
            t, jenis = ready, "ready_stock"
    # tahap aktif = kolom tahap yang terisi angka progress
    tahap = next((k for k in TAHAP if isinstance(t[k], (int, float))), None)
    return {"jenis": jenis, "spmk": t["spmk"], "lantai": t["lantai"], "progress": t["progress"], "tahap": tahap}


def main():
    ws = openpyxl.load_workbook(SRC, data_only=True, read_only=True)["LACAK"]
    proyek = clean(ws["C8"].value)
    records = []
    for row in ws.iter_rows(min_row=FIRST_ROW, max_col=column_index_from_string("DJ"), values_only=True):
        p = pick(row, PEMESANAN)
        # id = ID UNIT spreadsheet; Ready Stock (tanpa tgl order) & Cancel (tanpa ID UNIT) dilewati
        if not p["nama"] or not p["id_unit"] or not p["tgl_order"]:
            continue
        records.append({
            "id": p["id_unit"],
            "nik": NIK_DUMMY.get(p["blok"]),
            "proyek": proyek,
            "pemesanan": p,
            "teknik": teknik(row),
            "legal": pick(row, LEGAL),
            "keuangan": pick(row, KEUANGAN),
        })

    ids = [r["id"] for r in records]
    assert len(ids) == len(set(ids)), "ID UNIT duplikat"
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    print(f"{len(records)} konsumen -> {OUT.relative_to(ROOT)}")
    if records:  # jangan kosongkan tabel kalau Excel gagal terbaca
        sync_supabase(records)


def env(name):
    if name in os.environ:
        return os.environ[name]
    local = ROOT / ".env.local"
    for line in local.read_text().splitlines() if local.exists() else []:
        k, _, v = line.partition("=")
        if k.strip() == name:
            return v.strip().strip('"')
    return None


def sync_supabase(records):
    url = env("SUPABASE_URL")
    key = env("SUPABASE_SECRET_KEY") or env("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        sys.exit("SUPABASE_URL / SUPABASE_SECRET_KEY belum di-set di .env.local, data belum masuk Supabase")
    headers = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json"}

    def call(method, query, body=None, prefer="return=minimal"):
        req = urllib.request.Request(
            f"{url}/rest/v1/lacak_konsumen?{query}", method=method,
            data=json.dumps(body, ensure_ascii=False).encode() if body is not None else None,
            headers={**headers, "Prefer": prefer},
        )
        urllib.request.urlopen(req).close()

    rows = [{"id": r["id"], "nik": r["nik"], "data": r} for r in records]
    call("POST", "on_conflict=id", rows, "resolution=merge-duplicates,return=minimal")
    keep = ",".join('"' + r["id"].replace('"', '\\"') + '"' for r in records)
    call("DELETE", "id=not.in." + urllib.parse.quote(f"({keep})"))
    print(f"{len(rows)} konsumen -> Supabase lacak_konsumen")


if __name__ == "__main__":
    main()
